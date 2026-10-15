import os
import asyncio
import logging
import random
from datetime import datetime
from zoneinfo import ZoneInfo
 
import aiosqlite
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from pyrogram import Client, filters, idle, enums
from pyrogram.types import (
    Message, ChatMemberUpdated, ChatJoinRequest,
    ReplyKeyboardMarkup, KeyboardButton, ReplyKeyboardRemove,
)
from pyrogram.errors import (
    FloodWait, RPCError, ChatAdminRequired, UserAdminInvalid,
    PeerIdInvalid, ChannelPrivate, UserAlreadyParticipant,
)
from pyrogram.raw import functions, types as raw_types
 
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)
log = logging.getLogger("MultiControlBot")
 
# ==================== CONFIG ====================
API_ID = int(os.environ.get("API_ID", 0))
API_HASH = os.environ.get("API_HASH", "")
BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
OWNER_ID = int(os.environ.get("OWNER_ID", 0))
SESSION_STRING = os.environ.get("SESSION_STRING", "")
DB_PATH = os.environ.get("DB_PATH", "/data/bot.db")  # mount a Railway volume at /data
EXCEL_PATH = os.path.join(os.path.dirname(DB_PATH) or ".", "users.xlsx")
TZ = ZoneInfo(os.environ.get("TIMEZONE", "Asia/Kolkata"))
# Optional: only accept numbers starting with these country codes, e.g. "91" or "91,1". Empty = any.
ALLOWED_CODES = [c.strip().lstrip("+") for c in os.environ.get("ALLOWED_COUNTRY_CODES", "").split(",") if c.strip()]
 
if not all([API_ID, API_HASH, BOT_TOKEN, OWNER_ID]):
    raise SystemExit("Missing one of: API_ID, API_HASH, BOT_TOKEN, OWNER_ID")
if not SESSION_STRING:
    raise SystemExit("Missing SESSION_STRING (assistant account is required for VC control)")
 
# Bot account: receives commands, handles GBan
app = Client(
    "MultiControlBot",
    api_id=API_ID,
    api_hash=API_HASH,
    bot_token=BOT_TOKEN,
    in_memory=True,
)
 
# Assistant user account: performs voice chat actions (bots are not allowed to)
assistant = Client(
    "Assistant",
    api_id=API_ID,
    api_hash=API_HASH,
    session_string=SESSION_STRING,
    in_memory=True,
    no_updates=True,
)
 
# In-memory caches (backed by SQLite)
GBANNED: set[int] = set()
KNOWN_CHATS: set[int] = set()
BOT_ID: int = 0
ASSISTANT_ID: int = 0
 
ADMIN_STATUSES = (enums.ChatMemberStatus.ADMINISTRATOR, enums.ChatMemberStatus.OWNER)
 
 
# ==================== DATABASE ====================
async def db_init():
    os.makedirs(os.path.dirname(DB_PATH) or ".", exist_ok=True)
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("CREATE TABLE IF NOT EXISTS chats (chat_id INTEGER PRIMARY KEY)")
        await db.execute(
            "CREATE TABLE IF NOT EXISTS gbans ("
            "user_id INTEGER PRIMARY KEY, reason TEXT, "
            "banned_at TEXT DEFAULT CURRENT_TIMESTAMP)"
        )
        await db.execute(
            "CREATE TABLE IF NOT EXISTS users ("
            "user_id INTEGER PRIMARY KEY, username TEXT, phone TEXT, joined_at TEXT)"
        )
        await db.execute(
            "CREATE TABLE IF NOT EXISTS pending ("
            "user_id INTEGER, chat_id INTEGER, PRIMARY KEY (user_id, chat_id))"
        )
        await db.commit()
        async with db.execute("SELECT chat_id FROM chats") as cur:
            KNOWN_CHATS.update(r[0] for r in await cur.fetchall())
        async with db.execute("SELECT user_id FROM gbans") as cur:
            GBANNED.update(r[0] for r in await cur.fetchall())
    log.info("DB ready: %d chats, %d gbanned users", len(KNOWN_CHATS), len(GBANNED))
 
 
async def db_exec(query: str, params: tuple = ()):
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(query, params)
        await db.commit()
 
 
async def remember_chat(chat_id: int):
    if chat_id not in KNOWN_CHATS:
        KNOWN_CHATS.add(chat_id)
        await db_exec("INSERT OR IGNORE INTO chats (chat_id) VALUES (?)", (chat_id,))
 
 
async def forget_chat(chat_id: int):
    KNOWN_CHATS.discard(chat_id)
    await db_exec("DELETE FROM chats WHERE chat_id = ?", (chat_id,))
 
 
# ==================== HELPERS ====================
async def is_admin(client: Client, message: Message) -> bool:
    """True for owner, group admins, and anonymous admins."""
    if message.sender_chat and message.sender_chat.id == message.chat.id:
        return True  # anonymous admin
    user = message.from_user
    if not user:
        return False
    if user.id == OWNER_ID:
        return True
    try:
        member = await client.get_chat_member(message.chat.id, user.id)
        return member.status in ADMIN_STATUSES
    except RPCError:
        return False
 
 
async def assistant_can_manage_vc(chat_id: int) -> bool:
    try:
        me = await assistant.get_chat_member(chat_id, ASSISTANT_ID)
    except RPCError:
        return False
    if me.status == enums.ChatMemberStatus.OWNER:
        return True
    return bool(me.privileges and me.privileges.can_manage_video_chats)
 
 
async def get_group_call(chat_id: int):
    """Return the InputGroupCall of the active VC, or None (uses assistant account)."""
    await assistant.get_chat(chat_id)  # caches the peer for the assistant
    peer = await assistant.resolve_peer(chat_id)
    if isinstance(peer, raw_types.InputPeerChannel):
        full = await assistant.invoke(functions.channels.GetFullChannel(channel=peer))
    else:  # basic group
        full = await assistant.invoke(functions.messages.GetFullChat(chat_id=-chat_id))
    return getattr(full.full_chat, "call", None)
 
 
async def safe_ban(client: Client, chat_id: int, user_id: int) -> str:
    """Returns 'ok' or an error code string."""
    for _ in range(3):
        try:
            await client.ban_chat_member(chat_id, user_id)
            return "ok"
        except FloodWait as e:
            await asyncio.sleep(e.value + 1)
        except (ChatAdminRequired, UserAdminInvalid):
            return "no_rights_or_target_is_admin"
        except PeerIdInvalid:
            return "user_not_seen_by_bot"
        except ChannelPrivate:
            await forget_chat(chat_id)
            return "bot_removed"
        except RPCError as e:
            return type(e).__name__
    return "flood"
 
 
async def safe_unban(client: Client, chat_id: int, user_id: int) -> str:
    for _ in range(3):
        try:
            await client.unban_chat_member(chat_id, user_id)
            return "ok"
        except FloodWait as e:
            await asyncio.sleep(e.value + 1)
        except ChannelPrivate:
            await forget_chat(chat_id)
            return "bot_removed"
        except RPCError as e:
            return type(e).__name__
    return "flood"
 
 
async def parse_target(client: Client, message: Message):
    """Target from reply, numeric ID, or @username. Returns (user_id, reason)."""
    args = message.command[1:]
    if message.reply_to_message and message.reply_to_message.from_user:
        return message.reply_to_message.from_user.id, " ".join(args)
    if not args:
        return None, ""
    token, reason = args[0], " ".join(args[1:])
    try:
        return int(token), reason
    except ValueError:
        pass
    try:
        user = await client.get_users(token)
        return user.id, reason
    except RPCError:
        return None, ""
 
 
# ==================== 1. GLOBAL BAN ====================
owner_only = filters.user(OWNER_ID)
 
 
@app.on_message(filters.command("gban") & owner_only)
async def gban_cmd(client: Client, message: Message):
    target_id, reason = await parse_target(client, message)
    if not target_id:
        return await message.reply_text(
            "Usage: `/gban <user_id|@username> [reason]` or reply to a user with `/gban [reason]`"
        )
    if target_id in (OWNER_ID, BOT_ID, ASSISTANT_ID):
        return await message.reply_text("❌ I won't ban the owner or my own accounts.")
 
    reason = reason or "No reason provided."
    GBANNED.add(target_id)
    await db_exec(
        "INSERT OR REPLACE INTO gbans (user_id, reason) VALUES (?, ?)", (target_id, reason)
    )
 
    status = await message.reply_text(
        f"🔄 Banning `{target_id}` in {len(KNOWN_CHATS)} groups..."
    )
 
    ok, failures = 0, {}
    for chat_id in list(KNOWN_CHATS):
        result = await safe_ban(client, chat_id, target_id)
        if result == "ok":
            ok += 1
        else:
            failures[result] = failures.get(result, 0) + 1
        await asyncio.sleep(0.3)  # gentle rate limiting
 
    fail_text = "\n".join(f"   • {k}: {v}" for k, v in failures.items()) or "   • none"
    await status.edit_text(
        f"✅ **Global Ban Executed**\n\n"
        f"👤 **User ID:** `{target_id}`\n"
        f"📝 **Reason:** {reason}\n"
        f"🌐 **Banned in:** {ok} groups\n"
        f"⚠️ **Failed:**\n{fail_text}\n\n"
        f"ℹ️ The user is saved and will be auto-banned on join or on any message."
    )
 
 
@app.on_message(filters.command("ungban") & owner_only)
async def ungban_cmd(client: Client, message: Message):
    target_id, _ = await parse_target(client, message)
    if not target_id:
        return await message.reply_text("Usage: `/ungban <user_id|@username>` or reply to a user.")
 
    GBANNED.discard(target_id)
    await db_exec("DELETE FROM gbans WHERE user_id = ?", (target_id,))
 
    status = await message.reply_text(f"🔄 Unbanning `{target_id}`...")
    ok = 0
    for chat_id in list(KNOWN_CHATS):
        if await safe_unban(client, chat_id, target_id) == "ok":
            ok += 1
        await asyncio.sleep(0.3)
    await status.edit_text(f"✅ `{target_id}` removed from GBan list. Unbanned in {ok} groups.")
 
 
@app.on_message(filters.command("gbanlist") & owner_only)
async def gbanlist_cmd(client: Client, message: Message):
    if not GBANNED:
        return await message.reply_text("The GBan list is empty.")
    lines = [f"• `{uid}`" for uid in sorted(GBANNED)]
    await message.reply_text(f"🚫 **Globally banned users ({len(lines)}):**\n" + "\n".join(lines[:100]))
 
 
@app.on_message(filters.group, group=-1)
async def track_and_enforce(client: Client, message: Message):
    """Runs on every group message: remember the chat and enforce GBans."""
    await remember_chat(message.chat.id)
 
    offenders = []
    if message.new_chat_members:
        offenders += [u.id for u in message.new_chat_members if u.id in GBANNED]
    if message.from_user and message.from_user.id in GBANNED:
        offenders.append(message.from_user.id)
 
    for uid in set(offenders):
        result = await safe_ban(client, message.chat.id, uid)
        if result == "ok":
            log.info("Auto-banned gbanned user %s in %s", uid, message.chat.id)
            try:
                await message.delete()
            except RPCError:
                pass
            try:
                await client.send_message(
                    message.chat.id, f"🚨 Globally banned user `{uid}` was removed."
                )
            except RPCError:
                pass
        else:
            log.warning("Could not ban %s in %s: %s", uid, message.chat.id, result)
 
 
@app.on_chat_member_updated(filters.group)
async def bot_membership_tracker(client: Client, update: ChatMemberUpdated):
    """Track groups when the bot is added/removed."""
    new = update.new_chat_member
    if not new or not new.user.is_self:
        return
    if new.status in (enums.ChatMemberStatus.MEMBER, *ADMIN_STATUSES):
        await remember_chat(update.chat.id)
        log.info("Added to chat %s", update.chat.id)
    else:
        await forget_chat(update.chat.id)
        log.info("Removed from chat %s", update.chat.id)
 
 
# ==================== 2. VOICE CHAT START / END (via assistant) ====================
NEED_RIGHTS = (
    "❌ The assistant account must be in this group as admin with "
    "**Manage Video Chats** enabled."
)
 
 
@app.on_message(filters.command(["vcon", "startvc"]) & filters.group)
async def vc_on(client: Client, message: Message):
    if not await is_admin(client, message):
        return await message.reply_text("❌ Only group admins can control voice chats.")
    if not await assistant_can_manage_vc(message.chat.id):
        return await message.reply_text(NEED_RIGHTS)
    try:
        if await get_group_call(message.chat.id):
            return await message.reply_text("⚠️ A voice chat is already active.")
        await assistant.invoke(
            functions.phone.CreateGroupCall(
                peer=await assistant.resolve_peer(message.chat.id),
                random_id=random.randint(1, 2**31 - 1),
            )
        )
        await message.reply_text("🎙️ **Voice Chat** started!")
    except RPCError as e:
        log.exception("vcon failed")
        await message.reply_text(f"❌ Failed to start Voice Chat: `{e}`")
 
 
@app.on_message(filters.command(["vcoff", "endvc"]) & filters.group)
async def vc_off(client: Client, message: Message):
    if not await is_admin(client, message):
        return await message.reply_text("❌ Only group admins can control voice chats.")
    if not await assistant_can_manage_vc(message.chat.id):
        return await message.reply_text(NEED_RIGHTS)
    try:
        call = await get_group_call(message.chat.id)
        if not call:
            return await message.reply_text("⚠️ No active Voice Chat in this group.")
        await assistant.invoke(functions.phone.DiscardGroupCall(call=call))
        await message.reply_text("🛑 **Voice Chat** ended.")
    except RPCError as e:
        log.exception("vcoff failed")
        await message.reply_text(f"❌ Failed to end Voice Chat: `{e}`")
 
 
# ==================== 3. VC MUTE / UNMUTE (via assistant) ====================
async def _vc_set_mute(client: Client, message: Message, mute: bool):
    if not await is_admin(client, message):
        return await message.reply_text("❌ Admin command only.")
    if not message.reply_to_message or not message.reply_to_message.from_user:
        return await message.reply_text("⚠️ Reply to the user's message.")
    if not await assistant_can_manage_vc(message.chat.id):
        return await message.reply_text(NEED_RIGHTS)
 
    target = message.reply_to_message.from_user
    try:
        call = await get_group_call(message.chat.id)
        if not call:
            return await message.reply_text("⚠️ No active Voice Chat.")
        await assistant.get_chat_member(message.chat.id, target.id)  # cache peer
        await assistant.invoke(
            functions.phone.EditGroupCallParticipant(
                call=call,
                participant=await assistant.resolve_peer(target.id),
                muted=mute,
            )
        )
        if mute:
            await message.reply_text(f"🔇 Muted {target.mention} in the Voice Chat.")
        else:
            await message.reply_text(
                f"🔊 {target.mention} may speak again (they must unmute themselves)."
            )
    except RPCError as e:
        log.exception("vc mute/unmute failed")
        await message.reply_text(
            f"❌ Error: `{e}`\n(The user must currently be joined in the voice chat.)"
        )
 
 
@app.on_message(filters.command("vcmute") & filters.group)
async def vc_mute(client: Client, message: Message):
    await _vc_set_mute(client, message, True)
 
 
@app.on_message(filters.command("vcunmute") & filters.group)
async def vc_unmute(client: Client, message: Message):
    await _vc_set_mute(client, message, False)
 
 
 
# ==================== 4. JOIN REQUEST VERIFICATION + EXCEL ====================
EXCEL_LOCK = asyncio.Lock()
HEADERS = ["SNO", "USERNAME", "USER ID", "JOINED DATE", "MOBILE NUMBER"]
 
CONTACT_KB = ReplyKeyboardMarkup(
    [[KeyboardButton("📱 Share my mobile number", request_contact=True)]],
    resize_keyboard=True,
    one_time_keyboard=True,
)
 
 
def _write_excel(rows: list):
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for i, (user_id, username, phone, joined_at) in enumerate(rows, start=1):
        ws.append([i, username or "N/A", int(user_id), joined_at, phone])
        ws.cell(row=i + 1, column=5).number_format = "@"  # keep number as text
        ws.cell(row=i + 1, column=3).number_format = "0"
    for col, width in zip("ABCDE", (8, 24, 16, 22, 18)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(EXCEL_PATH)
 
 
async def export_excel() -> str:
    """Rebuild users.xlsx from the DB and return its path."""
    async with EXCEL_LOCK:
        async with aiosqlite.connect(DB_PATH) as db:
            async with db.execute(
                "SELECT user_id, username, phone, joined_at FROM users ORDER BY rowid"
            ) as cur:
                rows = await cur.fetchall()
        await asyncio.to_thread(_write_excel, rows)
    return EXCEL_PATH
 
 
async def get_verified_user(user_id: int):
    async with aiosqlite.connect(DB_PATH) as db:
        async with db.execute("SELECT phone FROM users WHERE user_id = ?", (user_id,)) as cur:
            return await cur.fetchone()
 
 
async def get_pending_chats(user_id: int) -> list:
    async with aiosqlite.connect(DB_PATH) as db:
        async with db.execute("SELECT chat_id FROM pending WHERE user_id = ?", (user_id,)) as cur:
            return [r[0] for r in await cur.fetchall()]
 
 
async def save_user(user, phone: str):
    joined = datetime.now(TZ).strftime("%d-%m-%Y %H:%M:%S")
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "INSERT INTO users (user_id, username, phone, joined_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(user_id) DO UPDATE SET username = excluded.username, phone = excluded.phone",
            (user.id, user.username, phone, joined),
        )
        await db.commit()
    await export_excel()
 
 
async def approve_pending(client: Client, user_id: int) -> int:
    """Approve every pending join request of this user. Returns how many were approved."""
    approved = 0
    for chat_id in await get_pending_chats(user_id):
        try:
            await client.approve_chat_join_request(chat_id, user_id)
            approved += 1
        except UserAlreadyParticipant:
            approved += 1
        except FloodWait as e:
            await asyncio.sleep(e.value + 1)
            try:
                await client.approve_chat_join_request(chat_id, user_id)
                approved += 1
            except RPCError:
                log.warning("Approve retry failed for %s in %s", user_id, chat_id)
        except ChannelPrivate:
            await forget_chat(chat_id)
        except RPCError as e:
            log.warning("Approve failed for %s in %s: %s", user_id, chat_id, e)
        await db_exec("DELETE FROM pending WHERE user_id = ? AND chat_id = ?", (user_id, chat_id))
    return approved
 
 
async def ask_for_number(client: Client, user_id: int, chat_title: str = "the group"):
    await client.send_message(
        user_id,
        f"👋 Hello! You requested to join **{chat_title}**.\n\n"
        "To verify your account, tap the button below to **share your mobile number**.\n"
        "Only the number linked to your own Telegram account is accepted.",
        reply_markup=CONTACT_KB,
    )
 
 
@app.on_chat_join_request()
async def on_join_request(client: Client, req: ChatJoinRequest):
    user, chat = req.from_user, req.chat
    log.info("JOIN REQUEST received: user=%s chat=%s", user.id, chat.id)
    await remember_chat(chat.id)
 
    # Globally banned users are declined straight away
    if user.id in GBANNED:
        try:
            await client.decline_chat_join_request(chat.id, user.id)
        except RPCError:
            pass
        return
 
    await db_exec("INSERT OR IGNORE INTO pending (user_id, chat_id) VALUES (?, ?)", (user.id, chat.id))
 
    # Already verified before -> approve instantly
    if await get_verified_user(user.id):
        n = await approve_pending(client, user.id)
        log.info("Auto-approved verified user %s (%d request(s))", user.id, n)
        return
 
    # Prefer Telegram's direct chat id (lets the bot DM users it has never talked to)
    dm_id = getattr(req, "user_chat_id", None) or user.id
    try:
        await ask_for_number(client, dm_id, chat.title or "the group")
        log.info("Contact request sent to %s", user.id)
    except Exception:
        # User can open the bot and press /start later; the request stays pending
        log.exception("Could not DM user %s", user.id)
 
 
@app.on_message(filters.contact & filters.private)
async def on_contact(client: Client, message: Message):
    user, contact = message.from_user, message.contact
    if not user:
        return
 
    pending = await get_pending_chats(user.id)
    if not pending:
        return await message.reply_text(
            "ℹ️ You have no pending join request. Send a join request to the group first.",
            reply_markup=ReplyKeyboardRemove(),
        )
 
    # The shared contact must belong to the sender's own account
    if contact.user_id != user.id:
        return await message.reply_text(
            "❌ That is not **your** number. Please use the button below to share your own contact.",
            reply_markup=CONTACT_KB,
        )
 
    phone = contact.phone_number.strip()
    if not phone.startswith("+"):
        phone = "+" + phone
 
    digits = phone.lstrip("+")
    if ALLOWED_CODES and not any(digits.startswith(c) for c in ALLOWED_CODES):
        return await message.reply_text(
            "❌ Your number's country code is not allowed for this group.",
            reply_markup=ReplyKeyboardRemove(),
        )
 
    await save_user(user, phone)
    n = await approve_pending(client, user.id)
    if n:
        await message.reply_text(
            "✅ Verified! Your join request has been **approved**. Welcome!",
            reply_markup=ReplyKeyboardRemove(),
        )
    else:
        await message.reply_text(
            "⚠️ Verified, but I couldn't approve the request. Make sure I'm admin with "
            "**Invite Users** permission, or ask an admin.",
            reply_markup=ReplyKeyboardRemove(),
        )
    log.info("Verified %s (%s), approved in %d chat(s)", user.id, phone, n)
 
 
@app.on_message(filters.private & filters.text & ~filters.command(["start", "list", "gban", "ungban", "gbanlist"]))
async def remind_contact(client: Client, message: Message):
    if message.from_user and await get_pending_chats(message.from_user.id):
        await message.reply_text("📱 Please tap the button to share your mobile number.", reply_markup=CONTACT_KB)
 
 
@app.on_message(filters.command("list") & filters.private & owner_only)
async def list_cmd(client: Client, message: Message):
    path = await export_excel()
    async with aiosqlite.connect(DB_PATH) as db:
        async with db.execute("SELECT COUNT(*) FROM users") as cur:
            total = (await cur.fetchone())[0]
    if total == 0:
        return await message.reply_text("The user list is empty.")
    await message.reply_document(
        path,
        file_name="users_list.xlsx",
        caption=f"📊 Verified users: **{total}**\n🕒 {datetime.now(TZ).strftime('%d-%m-%Y %H:%M:%S')}",
    )
 
 
# ==================== START ====================
@app.on_message(filters.command("start") & filters.private)
async def start_handler(client: Client, message: Message):
    user = message.from_user
    if user and await get_pending_chats(user.id) and not await get_verified_user(user.id):
        return await ask_for_number(client, user.id)
    await message.reply_text(
        "👋 **Multi-Control Bot Online!**\n\n"
        "**Owner:** `/gban`, `/ungban`, `/gbanlist`, `/list`\n"
        "**Admins (groups):** `/vcon`, `/vcoff`, `/vcmute`, `/vcunmute`"
    )
 
 
async def main():
    global BOT_ID, ASSISTANT_ID
    await db_init()
    await app.start()
    await assistant.start()
    BOT_ID = (await app.get_me()).id
    ASSISTANT_ID = (await assistant.get_me()).id
    log.info("Bot %s and assistant %s started", BOT_ID, ASSISTANT_ID)
    await idle()
    await assistant.stop()
    await app.stop()
 
 
if __name__ == "__main__":
    app.run(main())
